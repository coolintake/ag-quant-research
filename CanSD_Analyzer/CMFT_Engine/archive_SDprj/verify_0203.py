from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_Main.xlsx")
ws = wb['Wheat_Mo']
print('Dimensions:', ws.dimensions)
print('Max row:', ws.max_row, 'Max col:', ws.max_column)
print()
# Check rows for Food, Industrial, Food & Processing for 02-03
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=16, values_only=False):
    vals = [cell.value for cell in row[:16]]
    if vals[1] and ('Food' in str(vals[1]) or 'Industrial' in str(vals[1]) or 'Food & Processing' in str(vals[1])):
        if vals[2] == '02-03':
            print(f'Row {row[0].row}: {vals}')