from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_main.xlsx")
print('Sheets:', wb.sheetnames)
print()

# Check Wheat_Mo
ws_mo = wb['Wheat_Mo']
print('Wheat_Mo:')
print(f'  Dimensions: {ws_mo.dimensions}')
print(f'  Max row: {ws_mo.max_row}, Max col: {ws_mo.max_column}')
print(f'  Headers (row 2): {[ws_mo.cell(row=2, column=c).value for c in range(1, 17)]}')
print(f'  First few metric rows (col B):')
for r in range(3, min(15, ws_mo.max_row+1)):
    print(f'    Row {r}: A={ws_mo.cell(row=r, column=1).value}, B={ws_mo.cell(row=r, column=2).value}, C={ws_mo.cell(row=r, column=3).value}')
print()

# Check Wheat_Yr
ws_yr = wb['Wheat_Yr']
print('Wheat_Yr:')
print(f'  Dimensions: {ws_yr.dimensions}')
print(f'  Max row: {ws_yr.max_row}, Max col: {ws_yr.max_column}')
print(f'  Headers (row 2): {[ws_yr.cell(row=2, column=c).value for c in range(1, min(10, ws_yr.max_column+1))]}')
print(f'  First few metric rows (col B):')
for r in range(3, min(15, ws_yr.max_row+1)):
    print(f'    Row {r}: A={ws_yr.cell(row=r, column=1).value}, B={ws_yr.cell(row=r, column=2).value}, C={ws_yr.cell(row=r, column=3).value}')
print()

# Check formulas in Wheat_Yr
print('Sample formulas in Wheat_Yr:')
for row in ws_yr.iter_rows(min_row=3, max_row=min(10, ws_yr.max_row), max_col=min(10, ws_yr.max_column)):
    for cell in row:
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            print(f'  {cell.coordinate}: {cell.value}')
            break
    else:
        continue
    break

# Check CHECK row formula
print()
print('CHECK row formulas:')
for row in ws_yr.iter_rows(min_row=3, max_row=ws_yr.max_row, max_col=min(5, ws_yr.max_column)):
    for cell in row:
        if cell.value == 'CHECK':
            check_row = cell.row
            print(f'  CHECK row: {check_row}')
            for c in range(4, min(8, ws_yr.max_column+1)):
                c_cell = ws_yr.cell(row=check_row, column=c)
                if c_cell.value and isinstance(c_cell.value, str) and c_cell.value.startswith('='):
                    print(f'    {c_cell.coordinate}: {c_cell.value}')