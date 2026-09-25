from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_Main.xlsx", read_only=True)
print('Sheets:', wb.sheetnames)
print()

# Check each Batch 1 *_Mo tab
for name in ['Durum_Mo', 'Canola_Mo', 'Barley_Mo', 'Oats_Mo', 'Dry peas_Mo', 'Lentils_Mo']:
    ws = wb[name]
    print(f'{name}: {ws.max_row} rows x {ws.max_column} cols')
    # Check first few data rows
    for row in ws.iter_rows(min_row=3, max_row=5, max_col=5, values_only=True):
        print(f'  {row}')
    # Check if Column P has SUM formulas
    p_cell = ws.cell(row=3, column=16)
    print(f'  P3 formula: {p_cell.value}')
    print()

# Check if *_Yr tabs are preserved
for name in ['Durum_Yr', 'Canola_Yr', 'Barley_Yr', 'Oats_Yr', 'Dry peas_Yr', 'Lentils_Yr']:
    if name in wb.sheetnames:
        print(f'{name}: PRESERVED')
    else:
        print(f'{name}: MISSING')

wb.close()