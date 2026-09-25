from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_main.xlsx")

# Check Wheat_Mo data values
ws_mo = wb['Wheat_Mo']
print('Wheat_Mo sample data (row 6 = Beginning Stocks):')
for c in range(4, 10):
    cell = ws_mo.cell(row=6, column=c)
    print(f'  {cell.coordinate}: {cell.value} (fmt: {cell.number_format})')

print()
print('Wheat_Mo styling check (row 6):')
cell = ws_mo.cell(row=6, column=4)
print(f'  Font: {cell.font.name}, {cell.font.size}, bold={cell.font.bold}')
print(f'  Fill: {cell.fill.start_color.rgb if cell.fill.start_color else "none"}')
print(f'  Alignment: {cell.alignment.horizontal}')
print(f'  Border: {cell.border.left.style if cell.border.left else "none"}')

print()
print('Wheat_Yr styling check (row 6):')
ws_yr = wb['Wheat_Yr']
cell = ws_yr.cell(row=6, column=4)
print(f'  Font: {cell.font.name}, {cell.font.size}, bold={cell.font.bold}')
print(f'  Fill: {cell.fill.start_color.rgb if cell.fill.start_color else "none"}')
print(f'  Number format: {cell.number_format}')

# Check all commodities have both tabs
print()
print('All sheet pairs:')
for i in range(0, len(wb.sheetnames), 2):
    mo = wb.sheetnames[i]
    yr = wb.sheetnames[i+1] if i+1 < len(wb.sheetnames) else 'MISSING'
    print(f'  {mo} <-> {yr}')