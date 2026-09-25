import pandas as pd
from openpyxl import load_workbook

# Check CMFT_Master.csv structure
df = pd.read_csv(r"C:\Users\ahmed\OneDrive\Desktop\Python\CanSD_Analyzer\CMFT_Engine\data_store\CMFT_Master.csv")
print("CMFT_Master.csv columns:", list(df.columns))
print("Shape:", df.shape)
print()

# Check commodities
print("Unique Commodity_Norm:", sorted(df['Commodity_Norm'].unique()))
print()

# Check for Batch 1 commodities
batch1 = ['Durum', 'Canola', 'Barley', 'Oats', 'Dry peas', 'Lentils']
for c in batch1:
    subset = df[df['Commodity_Norm'] == c]
    print(f"{c}: {len(subset)} rows, Date range: {subset['Date_Str'].min()} to {subset['Date_Str'].max()}")
    if len(subset) > 0:
        print(f"  Metrics: {subset['Tier1_Commodity_Metric'].unique().tolist()}")
        print(f"  Months: {sorted(subset['Date'].dt.month.unique()) if 'Date' in df.columns else 'N/A'}")
print()

# Check CanadaSD_Main.xlsx sheets
wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_Main.xlsx", read_only=True)
print("CanadaSD_Main.xlsx sheets:", wb.sheetnames)
wb.close()