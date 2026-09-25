from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_main.xlsx")
print('Sheets:', wb.sheetnames)
print()

# Check Canola_Mo (Oilseed)
ws_mo = wb['Canola_Mo']
print('Canola_Mo:')
print(f'  Max row: {ws_mo.max_row}, Max col: {ws_mo.max_column}')
print(f'  First few metric rows (col B):')
for r in range(3, min(20, ws_mo.max_row+1)):
    print(f'    Row {r}: A={ws_mo.cell(row=r, column=1).value}, B={ws_mo.cell(row=r, column=2).value}, C={ws_mo.cell(row=r, column=3).value}')
print()

# Check Canola_Yr
ws_yr = wb['Canola_Yr']
print('Canola_Yr:')
print(f'  Max row: {ws_yr.max_row}, Max col: {ws_yr.max_column}')
print(f'  First few metric rows (col B):')
for r in range(3, min(20, ws_yr.max_row+1)):
    print(f'    Row {r}: A={ws_yr.cell(row=r, column=1).value}, B={ws_yr.cell(row=r, column=2).value}, C={ws_yr.cell(row=r, column=3).value}')
print()

# Check Soybeans_Mo (Oilseed, Sep-Aug)
ws_mo = wb['Soybeans_Mo']
print('Soybeans_Mo:')
print(f'  Max row: {ws_mo.max_row}, Max col: {ws_mo.max_column}')
print(f'  Headers (row 2): {[ws_mo.cell(row=2, column=c).value for c in range(1, 17)]}')
print(f'  First few metric rows (col B):')
for r in range(3, min(20, ws_mo.max_row+1)):
    print(f'    Row {r}: A={ws_mo.cell(row=r, column=1).value}, B={ws_mo.cell(row=r, column=2).value}, C={ws_mo.cell(row=r, column=3).value}')
print()

# Check Dry peas_Mo (Pulse)
ws_mo = wb['Dry peas_Mo']
print('Dry peas_Mo:')
print(f'  Max row: {ws_mo.max_row}, Max col: {ws_mo.max_column}')
print(f'  First few metric rows (col B):')
for r in range(3, min(20, ws_mo.max_row+1)):
    print(f'    Row {r}: A={ws_mo.cell(row=r, column=1).value}, B={ws_mo.cell(row=r, column=2).value}, C={ws_mo.cell(row=r, column=3).value}')
print()

# Check Dry peas_Yr
ws_yr = wb['Dry peas_Yr']
print('Dry peas_Yr:')
print(f'  Max row: {ws_yr.max_row}, Max col: {ws_yr.max_column}')
print(f'  First few metric rows (col B):')
for r in range(3, min(20, ws_yr.max_row+1)):
    print(f'    Row {r}: A={ws_yr.cell(row=r, column=1).value}, B={ws_yr.cell(row=r, column=2).value}, C={ws_yr.cell(row=r, column=3).value}')
print()

# Check formulas in Dry peas_Yr
print('Formulas in Dry peas_Yr:')
for row in ws_yr.iter_rows(min_row=3, max_row=ws_yr.max_row, max_col=ws_yr.max_column):
    for cell in row:
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            print(f'  {cell.coordinate}: {cell.value}')