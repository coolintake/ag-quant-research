import pandas as pd
from openpyxl import load_workbook

# Check CMFT_Master.csv structure
df = pd.read_csv(r"C:\Users\ahmed\OneDrive\Desktop\Python\CanSD_Analyzer\CMFT_Engine\data_store\CMFT_Master.csv")
print("CMFT_Master.csv columns:", list(df.columns))
print("Shape:", df.shape)
print()

# Check if there's a commodity column
print("First 5 rows:")
print(df.head().to_string())
print()

# Check crop years
if 'Crop_Year' in df.columns:
    print("Unique Crop_Year:", sorted(df['Crop_Year'].unique()))
if 'Month' in df.columns:
    print("Unique Months:", sorted(df['Month'].unique()))
if 'Crop_Year_Label' in df.columns:
    print("Unique Crop_Year_Label:", sorted(df['Crop_Year_Label'].unique()))
print()

# Check CanadaSD_Main.xlsx sheets
wb = load_workbook(r"C:\Users\ahmed\OneDrive\Desktop\Ahmed's Folder\OneKBCapital\Data\CanadaSD_Prj\CanadaSD_Main.xlsx", read_only=True)
print("CanadaSD_Main.xlsx sheets:", wb.sheetnames)
wb.close()