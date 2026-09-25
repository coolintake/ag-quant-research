"""
CGC_MonthlyEx.py
================
Monthly export S&D matrix: CGC Grain Statistics Weekly (GSW) weekly exports
pro-rated onto calendar months and laid out by marketing year.

    python CGC_MonthlyEx.py                 # refresh GSW cache, write workbook
    python CGC_MonthlyEx.py --no-refresh    # use gsw_data/_cache as-is

Output: Outputs/CGC_MonthlyEx_<latest week-ending date>.xlsx (next to this file).

Data logic
----------
* Weekly exports come from `cgc_engine.weekly_outflow` -- the same validated
  definition used everywhere else (Terminal Exports + Primary Shipment
  Distribution to Export Destinations, differenced from the Crop Year
  cumulative). Nothing about the export formula is re-implemented here.
* Each reporting week's volume is spread evenly over the days it covers
  (a regular week = weekly_exports / 7.0), then summed by calendar month.
  A week's span runs from the day after the previous reported week-end to
  its own week-ending date, so:
    - a regular Sunday-to-Sunday week is exactly 7 days (== /7.0);
    - a missing week makes the next diff cover 2+ weeks of volume, and it is
      spread over the matching 14+ days instead of being crammed into 7;
    - week 1 of a crop year starts no earlier than 1 August, so the CGC
      Aug-Jul crop-year total is never leaked into the prior July.
* Months are assigned to marketing years by commodity calendar:
  Aug-Jul (default), Sep-Aug (Soybeans), Oct-Sep (Corn). A marketing year
  whose first month precedes the first available data is dropped (it would
  otherwise show a misleadingly small partial total); elapsed-less future
  months of the current year are 0.
"""

from __future__ import annotations

import argparse
import logging
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Union

import numpy as np
import pandas as pd

from cgc_engine import (
    CROP_YEAR_SEASON_START_MONTH,
    DEFAULT_CROP_YEARS,
    DEFAULT_CURRENT_YEAR,
    SEASON_START_WEEK,
    crop_year_start_int,
    weekly_outflow,
)

logger = logging.getLogger("CGC_MonthlyEx")
PathLike = Union[str, Path]

HERE = Path(__file__).resolve().parent
GSW_DATA_DIR = HERE / "gsw_data"
OUTPUT_DIR = HERE / "Outputs"

# ═══════════════════════════════════════════════════════════════════════════
# CONFIG
# ═══════════════════════════════════════════════════════════════════════════

MONTH_ABBR: List[str] = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN",
                         "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]

# Marketing-year start month by raw GSW grain name (case-insensitive).
# Everything not listed uses the standard Aug-Jul grain/oilseed calendar.
DEFAULT_MY_START_MONTH: int = CROP_YEAR_SEASON_START_MONTH  # 8 = AUG
MY_START_MONTH_OVERRIDES: Dict[str, int] = {"soybeans": 9, "corn": 10}

# Raw GSW grain name -> label used in the report. Anything not listed is
# shown exactly as CGC reports it (e.g. 'Amber Durum', never 'Durum').
DISPLAY_NAMES: Dict[str, str] = {"Wheat": "Wheat, excluding durum"}

# Report order; any other commodity present in the data follows alphabetically.
COMMODITY_ORDER: List[str] = [
    "Wheat", "Amber Durum", "Barley", "Oats", "Canola", "Flaxseed",
    "Peas", "Lentils", "Soybeans", "Corn",
]

METRIC_LABEL = "Exports"
DAYS_PER_WEEK = 7


def my_start_month(grain: str) -> int:
    return MY_START_MONTH_OVERRIDES.get(str(grain).strip().lower(), DEFAULT_MY_START_MONTH)


def month_order(start_month: int) -> List[str]:
    return [MONTH_ABBR[(start_month - 1 + i) % 12] for i in range(12)]


def total_label(start_month: int) -> str:
    order = month_order(start_month)
    return f"{order[0]}/{order[-1]}"


def my_label(start_year: int) -> str:
    """2024 -> '24-25'."""
    return f"{start_year % 100:02d}-{(start_year + 1) % 100:02d}"


def display_name(grain: str) -> str:
    return DISPLAY_NAMES.get(grain, grain)


def _sort_key(grain: str):
    return (0, COMMODITY_ORDER.index(grain), "") if grain in COMMODITY_ORDER else (1, 0, grain.lower())


# ═══════════════════════════════════════════════════════════════════════════
# CORE LOGIC
# ═══════════════════════════════════════════════════════════════════════════

@dataclass
class MonthlyExportsResult:
    tables: Dict[str, pd.DataFrame]   # display name -> exact S&D layout (ints)
    monthly_long: pd.DataFrame        # unrounded monthly kt, one row per month
    coverage: pd.DataFrame            # per-commodity first/last day covered
    as_of: pd.Timestamp               # latest week-ending date in the data
    as_of_label: str                  # e.g. '2026-27 week 7'


def week_ending_calendar(gsw_df: pd.DataFrame) -> pd.DataFrame:
    """(crop_year, grain_week) -> parsed week_ending_date.

    GSW stores the date as a 'DD/MM/YYYY' string that nothing upstream
    parses; ISO strings / real datetimes (e.g. from a re-saved cache) are
    also accepted. Parsing is strictly day-first -- a month-first parse of
    '03/12/2023' would silently move a December week into March.
    """
    wk = gsw_df[["crop_year", "grain_week", "week_ending_date"]].drop_duplicates()
    col = wk["week_ending_date"]
    if pd.api.types.is_datetime64_any_dtype(col):
        parsed = pd.to_datetime(col)
    else:
        raw = col.astype(str).str.strip()
        parsed = pd.to_datetime(raw, format="%d/%m/%Y", errors="coerce")
        todo = parsed.isna()
        if todo.any():
            parsed[todo] = pd.to_datetime(raw[todo].str[:10], format="%Y-%m-%d", errors="coerce")
    wk = wk.assign(week_end=parsed.dt.normalize()).dropna(subset=["week_end", "grain_week"])

    n_dates = wk.groupby(["crop_year", "grain_week"])["week_end"].nunique()
    if (n_dates > 1).any():
        logger.warning("Multiple week_ending_date values for %d crop_year/week keys; using the latest.",
                       int((n_dates > 1).sum()))
    cal = wk.groupby(["crop_year", "grain_week"], as_index=False)["week_end"].max()

    # Sanity: every week must end inside its own crop year (Aug 1 .. next Aug 31).
    cy_start = pd.to_datetime(cal["crop_year"].map(crop_year_start_int).astype(str) + "-08-01")
    bad = (cal["week_end"] < cy_start) | (cal["week_end"] > cy_start + pd.DateOffset(years=1, months=1))
    if bad.any():
        logger.warning("%d week-ending dates fall outside their crop year -- check date format: %s",
                       int(bad.sum()), cal[bad].head(5).to_dict("records"))
    return cal


def _week_spans(weekly: pd.DataFrame) -> pd.DataFrame:
    """Add span_start / span_days to each (grain, week) row. Rows are kept
    even when their volume is NaN so the day chain stays intact."""
    w = weekly.sort_values(["grain", "week_end"]).reset_index(drop=True)
    cy_start = pd.to_datetime(w["crop_year"].map(crop_year_start_int).astype(str)
                              + f"-{CROP_YEAR_SEASON_START_MONTH:02d}-01")
    g = w.groupby("grain", sort=False)
    prev_end = g["week_end"].shift(1)
    prev_cy = g["crop_year"].shift(1)

    start = prev_end + pd.Timedelta(days=1)
    new_season = prev_cy.ne(w["crop_year"]) & prev_end.notna()
    start = start.where(~new_season, np.maximum(start, cy_start))

    first = prev_end.isna()
    first_default = np.where(w["grain_week"] <= SEASON_START_WEEK, cy_start,
                             w["week_end"] - pd.Timedelta(days=DAYS_PER_WEEK - 1))
    start = start.where(~first, pd.Series(first_default, index=w.index))
    start = np.minimum(start, w["week_end"])  # guard against duplicate dates

    w["span_start"] = pd.to_datetime(start)
    w["span_days"] = (w["week_end"] - w["span_start"]).dt.days.astype(int) + 1

    gaps = w[(w["span_days"] > DAYS_PER_WEEK) & ~new_season & ~first & w["weekly_outflow_ktonnes"].notna()]
    for (grain, cy), grp in gaps.groupby(["grain", "crop_year"]):
        logger.info("Missing reporting week(s) for %s %s: %d week(s) spread over their actual gap.",
                    grain, cy, len(grp))
    return w


def build_monthly_exports(
    gsw_df: pd.DataFrame,
    weekly: Optional[pd.DataFrame] = None,
    drop_truncated_leading_years: bool = True,
) -> MonthlyExportsResult:
    """Standardized GSW dataset -> monthly export S&D tables.

    `weekly` may be passed to reuse an already-computed
    `cgc_engine.weekly_outflow` result (e.g. `CGCAnalytics.outflow`).
    """
    if weekly is None:
        weekly = weekly_outflow(gsw_df)
    cal = week_ending_calendar(gsw_df)
    w = weekly.merge(cal, on=["crop_year", "grain_week"], how="left")
    if w["week_end"].isna().any():
        logger.warning("%d weekly rows have no week_ending_date and are skipped.", int(w["week_end"].isna().sum()))
        w = w.dropna(subset=["week_end"])

    w = _week_spans(w)
    as_of = w["week_end"].max()
    last_row = w.loc[w["week_end"] == as_of].iloc[0]
    as_of_label = f"{last_row['crop_year']} week {int(last_row['grain_week'])}"

    # Coverage is taken from the day chain (incl. NaN-volume weeks).
    coverage = (w.groupby("grain")
                 .agg(first_day=("span_start", "min"), last_day=("week_end", "max"))
                 .reset_index())

    v = w.dropna(subset=["weekly_outflow_ktonnes"])
    v = v[v["span_days"] > 0]

    # ---- expand each week into its days (vectorized) ----
    span = v["span_days"].to_numpy()
    idx = np.repeat(np.arange(len(v)), span)
    offset = np.arange(idx.size) - np.repeat(np.cumsum(span) - span, span)
    days = pd.DatetimeIndex(v["span_start"].to_numpy()[idx] + offset.astype("timedelta64[D]"))
    rate = (v["weekly_outflow_ktonnes"].to_numpy() / span)[idx]
    grains = v["grain"].to_numpy()[idx]

    daily = pd.DataFrame({"grain": grains, "month": days.to_period("M"), "kt": rate})
    monthly = daily.groupby(["grain", "month"], as_index=False)["kt"].sum()

    sm = monthly["grain"].map(my_start_month).to_numpy()
    mo = monthly["month"].dt.month.to_numpy()
    yr = monthly["month"].dt.year.to_numpy()
    monthly["my_start"] = np.where(mo >= sm, yr, yr - 1)
    monthly["pos"] = (mo - sm) % 12

    # ---- wide S&D tables ----
    tables: Dict[str, pd.DataFrame] = {}
    long_rows = []
    for grain in sorted(coverage["grain"], key=_sort_key):
        s_month = my_start_month(grain)
        cov = coverage.loc[coverage["grain"] == grain].iloc[0]
        first_my = cov["first_day"].year if cov["first_day"].month >= s_month else cov["first_day"].year - 1
        last_my = cov["last_day"].year if cov["last_day"].month >= s_month else cov["last_day"].year - 1
        if drop_truncated_leading_years and cov["first_day"] > pd.Timestamp(first_my, s_month, 1):
            first_my += 1
        years = list(range(first_my, last_my + 1))
        if not years:
            continue

        g = monthly[(monthly["grain"] == grain) & monthly["my_start"].between(first_my, last_my)]
        mat = np.zeros((len(years), 12))
        np.add.at(mat, (g["my_start"].to_numpy() - first_my, g["pos"].to_numpy()), g["kt"].to_numpy())

        cols = month_order(s_month)
        ints = np.rint(mat).astype(np.int64)
        tbl = pd.DataFrame(ints, columns=cols)
        tbl.insert(0, "Crop Year", [my_label(y) for y in years])
        tbl.insert(0, "Metric", METRIC_LABEL)
        tbl.insert(0, "Commodity", display_name(grain))
        tbl[total_label(s_month)] = ints.sum(axis=1)
        tables[display_name(grain)] = tbl

        for i, y in enumerate(years):
            for p in range(12):
                long_rows.append((display_name(grain), total_label(s_month), my_label(y), cols[p],
                                  pd.Timestamp(y + (s_month - 1 + p) // 12, (s_month - 1 + p) % 12 + 1, 1),
                                  mat[i, p]))

    monthly_long = pd.DataFrame(long_rows, columns=["Commodity", "Calendar", "Crop Year", "Month",
                                                    "Month Start", "Exports (kt)"])
    coverage["Commodity"] = coverage["grain"].map(display_name)
    return MonthlyExportsResult(tables, monthly_long, coverage, as_of, as_of_label)


# ═══════════════════════════════════════════════════════════════════════════
# EXCEL OUTPUT
# ═══════════════════════════════════════════════════════════════════════════

def write_workbook(result: MonthlyExportsResult, output_dir: PathLike = OUTPUT_DIR,
                   filename: Optional[str] = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (filename or f"CGC_MonthlyEx_{result.as_of:%Y-%m-%d}.xlsx")

    F = "Arial"
    title_font = Font(name=F, size=13, bold=True, color="14532D")
    sub_font = Font(name=F, size=9, italic=True, color="555555")
    hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="14532D")
    body_font = Font(name=F, size=10)
    total_font = Font(name=F, size=10, bold=True)
    future_font = Font(name=F, size=10, color="B0B0B0")
    partial_fill = PatternFill("solid", fgColor="FFF4CC")
    total_fill = PatternFill("solid", fgColor="E8F3EC")
    thin = Side(style="thin", color="C8C8C8")
    num_fmt = "#,##0"

    as_of = result.as_of
    cov = result.coverage.set_index("Commodity")

    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Exports"
    ws["A1"] = "Canadian Grain Commission - Monthly Exports by Marketing Year (thousand tonnes)"
    ws["A1"].font = title_font
    ws["A2"] = (f"Source: CGC Grain Statistics Weekly. Data through week ending {as_of:%d %b %Y} "
                f"({result.as_of_label}). Weekly volumes pro-rated to calendar months by day.")
    ws["A2"].font = sub_font
    ws["A3"] = "Shaded yellow = month in progress (partial); grey = month not yet elapsed."
    ws["A3"].font = sub_font

    r = 5
    for name, tbl in result.tables.items():
        headers = list(tbl.columns)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(horizontal="left" if c <= 3 else "center")
        r += 1
        s_month = MONTH_ABBR.index(headers[3]) + 1
        last_day = cov.loc[name, "last_day"]
        for _, row in tbl.iterrows():
            y0 = 2000 + int(row["Crop Year"][:2])
            for c in range(1, 4):
                ws.cell(row=r, column=c, value=row.iloc[c - 1]).font = body_font
            for p in range(12):
                cell = ws.cell(row=r, column=4 + p, value=int(row.iloc[3 + p]))
                cell.number_format = num_fmt
                m_start = pd.Timestamp(y0 + (s_month - 1 + p) // 12, (s_month - 1 + p) % 12 + 1, 1)
                m_end = m_start + pd.offsets.MonthEnd(0)
                if m_start > last_day:
                    cell.font = future_font
                elif m_end > last_day:
                    cell.font, cell.fill = body_font, partial_fill
                else:
                    cell.font = body_font
            tc = ws.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})")
            tc.number_format, tc.font, tc.fill = num_fmt, total_font, total_fill
            for c in range(1, 17):
                ws.cell(row=r, column=c).border = Border(bottom=thin)
            r += 1
        r += 1  # blank spacer between commodity blocks

    widths = [26, 9, 10] + [8] * 12 + [11]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "D5"
    ws.sheet_view.showGridLines = False

    # --- long, unrounded data for pivots / charts ---
    wl = wb.create_sheet("Monthly Long")
    long_df = result.monthly_long
    for c, h in enumerate(long_df.columns, start=1):
        cell = wl.cell(row=1, column=c, value=h)
        cell.font, cell.fill = hdr_font, hdr_fill
    for i, rec in enumerate(long_df.itertuples(index=False), start=2):
        for c, val in enumerate(rec, start=1):
            cell = wl.cell(row=i, column=c, value=val.to_pydatetime() if isinstance(val, pd.Timestamp) else val)
            cell.font = body_font
            if c == 5:
                cell.number_format = "mmm-yyyy"
            elif c == 6:
                cell.number_format = "#,##0.000"
    for i, wdt in enumerate([26, 10, 10, 8, 12, 13], start=1):
        wl.column_dimensions[get_column_letter(i)].width = wdt
    wl.freeze_panes = "A2"
    wl.auto_filter.ref = f"A1:F{len(long_df) + 1}"

    # --- notes ---
    wn = wb.create_sheet("Notes")
    notes = [
        ("Units", "Thousand tonnes (kt). Matrix cells are rounded to integers; totals are the SUM of the rounded months."),
        ("Export definition", "cgc_engine.weekly_outflow: Terminal Exports (metric 'Exports') + Primary Shipment "
                              "Distribution to 'Export Destinations', Crop Year cumulative differenced week over week."),
        ("Pro-rating", "Each week's volume is spread evenly across the days it covers (regular week = weekly / 7). "
                       "A missing week's volume is spread across the full gap; week 1 starts no earlier than 1 Aug."),
        ("Calendars", "Aug-Jul for all commodities except Soybeans (Sep-Aug) and Corn (Oct-Sep)."),
        ("Crop Year label", "Marketing year by calendar date, e.g. Corn 25-26 = Oct 2025 - Sep 2026. For Soybeans/Corn "
                            "this differs from the CGC Aug-Jul crop-year cumulative."),
        ("Truncated years", "A marketing year that starts before the first available data is omitted."),
        ("Negative revisions", "weekly_outflow floors weekly values at 0, so a downward CGC revision is not netted out."),
        ("Data through", f"Week ending {as_of:%Y-%m-%d} ({result.as_of_label})."),
    ]
    wn["A1"], wn["A1"].font = "Methodology", title_font
    for i, (k, val) in enumerate(notes, start=3):
        wn.cell(row=i, column=1, value=k).font = total_font
        c = wn.cell(row=i, column=2, value=val)
        c.font, c.alignment = body_font, Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 20
    wn.column_dimensions["B"].width = 110

    wb.calculation.fullCalcOnLoad = True
    try:
        wb.save(path)
    except PermissionError as exc:
        raise SystemExit(f"Cannot write {path} -- is it open in Excel? Close it and re-run.") from exc
    return path


# ═══════════════════════════════════════════════════════════════════════════
# RUNNER
# ═══════════════════════════════════════════════════════════════════════════

def run(gsw_data_dir: PathLike = GSW_DATA_DIR, output_dir: PathLike = OUTPUT_DIR,
        refresh: bool = True, force_refresh: bool = False,
        keep_truncated: bool = False) -> Path:
    t0 = time.perf_counter()
    if refresh:
        from ingestion import CGCDownloader
        gsw = CGCDownloader(gsw_data_dir, years=DEFAULT_CROP_YEARS,
                            current_year=DEFAULT_CURRENT_YEAR).load_all(force_refresh=force_refresh)
    else:
        cache = Path(gsw_data_dir) / "_cache" / "combined_gsw.parquet"
        if not cache.exists():
            raise SystemExit(f"No cached GSW data at {cache}; run without --no-refresh first.")
        gsw = pd.read_parquet(cache)

    result = build_monthly_exports(gsw, drop_truncated_leading_years=not keep_truncated)
    path = write_workbook(result, output_dir)

    with pd.option_context("display.width", 250, "display.max_columns", 20):
        for tbl in result.tables.values():
            print(tbl.to_string(index=False), "\n")
    print(f"Data through {result.as_of:%Y-%m-%d} ({result.as_of_label}). "
          f"{len(result.tables)} commodities. Wrote {path} in {time.perf_counter() - t0:.1f}s")
    return path


def main(argv: Optional[List[str]] = None) -> None:
    ap = argparse.ArgumentParser(description="CGC monthly export S&D matrix -> Excel")
    ap.add_argument("--gsw-dir", default=str(GSW_DATA_DIR))
    ap.add_argument("--output-dir", default=str(OUTPUT_DIR))
    ap.add_argument("--no-refresh", action="store_true", help="use the cached parquet, no download")
    ap.add_argument("--force-refresh", action="store_true", help="re-download every crop year")
    ap.add_argument("--keep-truncated", action="store_true",
                    help="keep leading marketing years that start before the data does")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    run(a.gsw_dir, a.output_dir, refresh=not a.no_refresh,
        force_refresh=a.force_refresh, keep_truncated=a.keep_truncated)


if __name__ == "__main__":
    main(sys.argv[1:])
